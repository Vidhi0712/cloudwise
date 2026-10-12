import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

OUTPUT_FILE = "CloudWise_Project_Presentation.pptx"
SCREENSHOT_DIR = "screenshots"

# Color Palette (Dark Theme / Cloud FinOps SaaS)
BG_COLOR = RGBColor(11, 17, 33)        # #0B1121 Deep Navy
CARD_BG = RGBColor(19, 29, 53)         # #131D35 Card Background
CARD_BORDER = RGBColor(38, 54, 88)     # #263658 Card Stroke
PRIMARY = RGBColor(99, 102, 241)       # #6366F1 Modern Indigo Accent
CYAN_ACCENT = RGBColor(56, 189, 248)   # #38BDF8 Bright Cyan
SUCCESS = RGBColor(16, 185, 129)       # #10B981 Emerald
WARNING = RGBColor(245, 158, 11)       # #F59E0B Amber
DANGER = RGBColor(239, 68, 68)         # #EF4444 Red
TEXT_LIGHT = RGBColor(248, 250, 252)   # #F8FAFC White
TEXT_MUTED = RGBColor(148, 163, 184)   # #94A3B8 Slate Gray
TEXT_DIM = RGBColor(100, 116, 139)     # #64748B Dim Slate


def set_slide_background(slide):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = BG_COLOR


def add_header(slide, title_text, category="CLOUDWISE FINOPS PLATFORM", slide_num=None):
    set_slide_background(slide)

    # Category Pill / Badge
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
    tf = cat_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = category.upper()
    p.font.name = "Arial"
    p.font.size = Pt(9.5)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    # Main Title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(10), Inches(0.6))
    tf2 = title_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = title_text
    p2.font.name = "Arial"
    p2.font.size = Pt(22)
    p2.font.bold = True
    p2.font.color.rgb = TEXT_LIGHT

    # Accent Line under header
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.35), Inches(11.733), Inches(0.02))
    line.fill.solid()
    line.fill.fore_color.rgb = CARD_BORDER
    line.line.color.rgb = CARD_BORDER

    # Footer
    footer = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.3))
    tff = footer.text_frame
    tff.margin_left = tff.margin_top = tff.margin_right = tff.margin_bottom = 0
    pf = tff.paragraphs[0]
    pf.text = "CloudWise | Advanced Cloud Computing (BECE355L) | Pragya Prakash (24BML0075) & Vidhi Agarwal (24BML0081)"
    pf.font.name = "Arial"
    pf.font.size = Pt(8.5)
    pf.font.color.rgb = TEXT_DIM

    if slide_num is not None:
        pfn = tff.add_paragraph()
        pfn.text = f"Slide {slide_num} of 16"
        pfn.alignment = PP_ALIGN.RIGHT
        pfn.font.name = "Arial"
        pfn.font.size = Pt(8.5)
        pfn.font.color.rgb = TEXT_DIM


def create_card(slide, left, top, width, height, title=None, border_color=CARD_BORDER, bg_color=CARD_BG):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = bg_color
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1)

    if title:
        tb = slide.shapes.add_textbox(left + Inches(0.15), top + Inches(0.12), width - Inches(0.3), Inches(0.35))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title.upper()
        p.font.name = "Arial"
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = CYAN_ACCENT

    return shape


def add_bullet(tf, bold_prefix, text_content, font_size=11, color=TEXT_LIGHT, level=0):
    p = tf.add_paragraph() if len(tf.paragraphs[0].text) > 0 else tf.paragraphs[0]
    p.level = level
    p.space_after = Pt(4)

    run1 = p.add_run()
    run1.text = bold_prefix + " " if bold_prefix else ""
    run1.font.name = "Arial"
    run1.font.size = Pt(font_size)
    run1.font.bold = True
    run1.font.color.rgb = color

    run2 = p.add_run()
    run2.text = text_content
    run2.font.name = "Arial"
    run2.font.size = Pt(font_size)
    run2.font.color.rgb = TEXT_LIGHT


def build_all_slides():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    print("Building Slide 1: Title...")
    # ==========================================
    # SLIDE 1 — TITLE
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    # Decorative Backdrop Card
    create_card(s1, Inches(1.0), Inches(0.9), Inches(11.333), Inches(5.7), border_color=PRIMARY)

    # Top Tag
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.333), Inches(0.4))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "ACADEMIC PROJECT PRESENTATION | B.TECH FINAL EVALUATION"
    p.font.name = "Arial"
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = CYAN_ACCENT

    # Main Project Title
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.75), Inches(10.333), Inches(1.0))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "CloudWise"
    p.font.name = "Arial"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.color.rgb = TEXT_LIGHT

    # Main Subtitle
    tb = s1.shapes.add_textbox(Inches(1.5), Inches(2.75), Inches(10.333), Inches(0.8))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Explainable Cloud Cost Intelligence & Optimization Platform"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = PRIMARY

    p2 = tf.add_paragraph()
    p2.text = "AI-Assisted Cloud Cost Anomaly Detection and Resource Optimization"
    p2.font.name = "Arial"
    p2.font.size = Pt(14)
    p2.font.color.rgb = TEXT_MUTED

    # Divider line
    div = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(3.75), Inches(10.333), Inches(0.02))
    div.fill.solid()
    div.fill.fore_color.rgb = CARD_BORDER
    div.line.color.rgb = CARD_BORDER

    # Details Grid: Left = Team, Right = Course & Faculty
    # Left Card: Student Team
    create_card(s1, Inches(1.5), Inches(4.05), Inches(5.0), Inches(2.1), title="Project Team Members")
    tb = s1.shapes.add_textbox(Inches(1.7), Inches(4.5), Inches(4.6), Inches(1.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "• Pragya Prakash"
    p.font.bold = True
    p.font.size = Pt(13)
    p.font.color.rgb = TEXT_LIGHT
    p_sub = tf.add_paragraph()
    p_sub.text = "   Register No: 24BML0075"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = CYAN_ACCENT

    p2 = tf.add_paragraph()
    p2.text = "• Vidhi Agarwal"
    p2.font.bold = True
    p2.font.size = Pt(13)
    p2.font.color.rgb = TEXT_LIGHT
    p2_sub = tf.add_paragraph()
    p2_sub.text = "   Register No: 24BML0081"
    p2_sub.font.size = Pt(11)
    p2_sub.font.color.rgb = CYAN_ACCENT

    # Right Card: Academic Course & Guide
    create_card(s1, Inches(6.8), Inches(4.05), Inches(5.0), Inches(2.1), title="Course & Evaluation Guide")
    tb = s1.shapes.add_textbox(Inches(7.0), Inches(4.5), Inches(4.6), Inches(1.5))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    p.text = "Subject: Advanced Cloud Computing"
    p.font.bold = True
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_LIGHT

    p2 = tf.add_paragraph()
    p2.text = "Course Code: BECE355L"
    p2.font.size = Pt(11.5)
    p2.font.color.rgb = TEXT_MUTED

    p3 = tf.add_paragraph()
    p3.text = "Faculty / Evaluator: Prof. Sangeetha N"
    p3.font.bold = True
    p3.font.size = Pt(12)
    p3.font.color.rgb = SUCCESS

    p4 = tf.add_paragraph()
    p4.text = "Implementation Status: Fully Runnable & Verified MVP"
    p4.font.size = Pt(10.5)
    p4.font.color.rgb = CYAN_ACCENT


    print("Building Slide 2: Contents...")
    # ==========================================
    # SLIDE 2 — CONTENTS
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Table of Contents", "AGENDA", 2)

    topics = [
        ("01", "Objective", "Problem formulation, cloud FinOps motivations, and core platform targets."),
        ("02", "Proposed Solution", "Single-pane CloudWise architecture with simulated AWS data pipeline."),
        ("03", "Methodology", "7-step operational pipeline from telemetry generation to explainable savings."),
        ("04", "Tech Stack", "React 19, FastAPI, SQLite, Scikit-learn, and simulated AWS service concepts."),
        ("05", "Novelty", "Why CloudWise is novel: Explainable recommendations, confidence scores, and risk."),
        ("06", "Scalability & Future Scope", "Evolution path from reproducible MVP to enterprise multi-cloud AWS APIs."),
        ("07", "Outcomes", "Verified achievements: ₹16,310/mo savings, 61 anomalies, 8 recommendations."),
        ("08", "References", "Formal academic citations and official AWS, Scikit-learn, and FastAPI papers.")
    ]

    for i, (num, title, desc) in enumerate(topics):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.6 + row * 1.3)
        width = Inches(5.75)
        height = Inches(1.15)

        create_card(s2, left, top, width, height)

        # Number pill
        num_box = s2.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), Inches(0.7), Inches(0.7))
        ntf = num_box.text_frame
        np = ntf.paragraphs[0]
        np.text = num
        np.font.name = "Arial"
        np.font.size = Pt(22)
        np.font.bold = True
        np.font.color.rgb = PRIMARY

        # Title & desc
        text_box = s2.shapes.add_textbox(left + Inches(0.95), top + Inches(0.18), width - Inches(1.1), Inches(0.8))
        ttf = text_box.text_frame
        ttf.word_wrap = True
        tp1 = ttf.paragraphs[0]
        tp1.text = title
        tp1.font.name = "Arial"
        tp1.font.size = Pt(13)
        tp1.font.bold = True
        tp1.font.color.rgb = TEXT_LIGHT

        tp2 = ttf.add_paragraph()
        tp2.text = desc
        tp2.font.name = "Arial"
        tp2.font.size = Pt(9.5)
        tp2.font.color.rgb = TEXT_MUTED


    print("Building Slide 3: Objective...")
    # ==========================================
    # SLIDE 3 — OBJECTIVE
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Problem Statement & Project Objectives", "PROBLEM & GOALS", 3)

    # Left: Problem Context
    create_card(s3, Inches(0.8), Inches(1.6), Inches(4.6), Inches(5.1), title="The Problem in Modern Cloud Computing")
    tb = s3.shapes.add_textbox(Inches(1.0), Inches(2.1), Inches(4.2), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "Opaque Cloud Billing:", "Cloud infrastructure generates massive, continuously changing cost & usage streams. Standard dashboards only report aggregate bills.", font_size=11)
    add_bullet(tf, "Lack of Transparency:", "Native billing reports fail to explain:", font_size=11)
    add_bullet(tf, "•", "Why a sudden cost surge or traffic spike occurred.", font_size=10, level=1)
    add_bullet(tf, "•", "Which instances are idle or severely oversized.", font_size=10, level=1)
    add_bullet(tf, "•", "What specific engineering actions can safely lower bills.", font_size=10, level=1)
    add_bullet(tf, "•", "Exact quantifiable savings in local currency (₹ INR).", font_size=10, level=1)

    # Mini Flow Box inside Left
    m_box = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(5.35), Inches(4.2), Inches(1.15))
    m_box.fill.solid()
    m_box.fill.fore_color.rgb = RGBColor(15, 23, 42)
    m_box.line.color.rgb = CARD_BORDER
    mtf = m_box.text_frame
    mtf.word_wrap = True
    mp = mtf.paragraphs[0]
    mp.text = "Cloud Data → Analysis → Intelligence → Recommendation → Savings"
    mp.alignment = PP_ALIGN.CENTER
    mp.font.size = Pt(9.5)
    mp.font.bold = True
    mp.font.color.rgb = CYAN_ACCENT

    # Right: 7 Objectives of CloudWise
    create_card(s3, Inches(5.65), Inches(1.6), Inches(6.88), Inches(5.1), title="Core Objectives of CloudWise")
    tb2 = s3.shapes.add_textbox(Inches(5.85), Inches(2.1), Inches(6.48), Inches(4.4))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    objs = [
        ("1. Real-time Telemetry & Cost Monitoring:", "Track simulated cloud resources across compute, database, storage, and serverless tiers."),
        ("2. Machine Learning Anomaly Detection:", "Unsupervised detection of unusual daily spend spikes using Scikit-learn's Isolation Forest."),
        ("3. Underutilization & Inefficiency Profiling:", "Transparent FinOps heuristics identifying underused compute (<10% CPU), idle DBs, and cold storage."),
        ("4. Actionable Recommendation Generation:", "Generate concrete, non-destructive rightsizing and migration recommendations."),
        ("5. Quantified Savings Estimation:", "Compute exact monthly and annual savings in Indian Rupees (₹ INR) for every identified waste."),
        ("6. Explainable Human-Readable Rationale:", "Answer WHAT is wrong, WHY it was flagged, WHAT to do, RISK tier, and CONFIDENCE score."),
        ("7. Short-Term Predictive Cost Forecasting:", "Project next 14 days of cloud expenditure using regression modeling to prevent bill shocks.")
    ]
    for pref, text in objs:
        add_bullet(tf2, pref, text, font_size=10.5)


    print("Building Slide 4: Proposed Solution...")
    # ==========================================
    # SLIDE 4 — PROPOSED SOLUTION — ARCHITECTURE
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Proposed Solution — High-Level Architecture", "SYSTEM ARCHITECTURE", 4)

    # Quote banner
    qb = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.65))
    qb.fill.solid()
    qb.fill.fore_color.rgb = RGBColor(30, 41, 59)
    qb.line.color.rgb = PRIMARY
    qtf = qb.text_frame
    qp = qtf.paragraphs[0]
    qp.text = '"An explainable cloud cost intelligence platform that combines anomaly detection, resource utilization analysis, cost forecasting, and recommendation generation in a single unified dashboard."'
    qp.alignment = PP_ALIGN.CENTER
    qp.font.size = Pt(11)
    qp.font.italic = True
    qp.font.color.rgb = CYAN_ACCENT

    # Architecture Flow Boxes (5 Layers)
    layers = [
        ("LAYER 1: DATA INGESTION", "Simulated AWS Telemetry Dataset\n• 60-Day Historical Data | 20 Active Resources\n• Services: EC2, RDS, S3, Lambda\n• Injected Spikes, Underutilization & Traffic Bursts", PRIMARY),
        ("LAYER 2: PERSISTENCE", "SQLite Relational Storage (cloudwise.db)\n• Tables: resources, cost_records, usage_records\n• 1,200 Cost Records + 1,200 Metric Telemetry Points\n• Zero external database dependency, instant processing", RGBColor(51, 65, 85)),
        ("LAYER 3: ANALYTICS & ML ENGINE", "CloudWise Intelligence Core (Python / Scikit-Learn)\n┌──────────────────────┬──────────────────────┬──────────────────────┐\n│ ML Anomaly Detection │ Utilization Analysis │  14-Day Forecasting  │\n│   Isolation Forest   │  FinOps Heuristics   │  Linear Trend Model  │\n└──────────────────────┴──────────────────────┴──────────────────────┘", PRIMARY),
        ("LAYER 4: EXPLAINABILITY & INSIGHTS", "Dynamic Reasoning & Decision Layer\n• Evaluates confidence % and risk tier per resource\n• Dynamic calculation-backed explanations (No LLM tokens/costs)\n• CloudWise Insights Engine generating 4 high-value executive takeaways", SUCCESS),
        ("LAYER 5: PRESENTATION", "Modern SaaS Web Dashboard (React 19 + Recharts + Tailwind CSS)\n• Executive Dashboard  • Resource Intelligence  • Cost Deep-Dive  • Anomalies  • Recommendations", CYAN_ACCENT)
    ]

    for idx, (ltitle, ldesc, lcolor) in enumerate(layers):
        top_pos = Inches(2.3 + idx * 0.85)
        card = create_card(s4, Inches(0.8), top_pos, Inches(11.733), Inches(0.75), border_color=lcolor)
        tb = s4.shapes.add_textbox(Inches(1.0), top_pos + Inches(0.08), Inches(11.333), Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = ltitle
        p1.font.name = "Arial"
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = lcolor

        p2 = tf.add_paragraph()
        p2.text = ldesc
        p2.font.name = "Arial"
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_LIGHT

    # Bottom Note
    note_box = s4.shapes.add_textbox(Inches(0.8), Inches(6.6), Inches(11.733), Inches(0.35))
    ntf = note_box.text_frame
    np = ntf.paragraphs[0]
    np.text = "Note: CloudWise currently operates in Demo Mode using reproducible simulated AWS Cost Explorer/CloudWatch data. Live AWS API integration is planned for future scope."
    np.font.size = Pt(9)
    np.font.italic = True
    np.font.color.rgb = WARNING


    print("Building Slide 5: Methodology...")
    # ==========================================
    # SLIDE 5 — METHODOLOGY
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Methodology — 7-Step Operational Pipeline", "PROCESSING PIPELINE", 5)

    steps = [
        ("STEP 1", "Data Generation", "Synthesize 60 days of telemetry across 20 resources (EC2, RDS, S3, Lambda) with seed 42 to ensure reproducible anomalies and idle patterns."),
        ("STEP 2", "Data Storage", "Persist schemas in SQLite (`resources`, `cost_records`, `usage_records`) with automated initialization and seeding on startup."),
        ("STEP 3", "Data Preprocessing", "Extract rolling means, standard deviations, and service-level cost aggregations via Pandas and NumPy."),
        ("STEP 4", "AI / ML Detection", "Apply Scikit-learn's Isolation Forest (`contamination=0.08`, `n_estimators=40`) to isolate abnormal spend deviations."),
        ("STEP 5", "Utilization Heuristics", "Execute FinOps rules: EC2 CPU < 10%, RDS CPU < 15% & connections < 10, S3 access < 15 requests/day, Lambda burst detection."),
        ("STEP 6", "Explainability Engine", "Derive transparent metrics: Problem, Recommended Action, Estimated Savings (₹), Risk Level, Confidence %, and Reason."),
        ("STEP 7", "Visualization Dashboard", "Render interactive views in React with Recharts (Area, Donut, Stacked, Bar charts), responsive sidebar, and slide-out inspector.")
    ]

    for idx, (snum, stitle, sdesc) in enumerate(steps):
        # 2 columns layout: 4 on left, 3 on right
        if idx < 4:
            left = Inches(0.8)
            top = Inches(1.6 + idx * 1.25)
            width = Inches(5.7)
        else:
            left = Inches(6.8)
            top = Inches(1.6 + (idx - 4) * 1.65)
            width = Inches(5.7)

        h = Inches(1.15) if idx < 4 else Inches(1.5)
        create_card(s5, left, top, width, h, border_color=PRIMARY)

        tb = s5.shapes.add_textbox(left + Inches(0.2), top + Inches(0.12), width - Inches(0.4), h - Inches(0.24))
        tf = tb.text_frame
        tf.word_wrap = True

        p1 = tf.paragraphs[0]
        p1.text = f"{snum}: {stitle.upper()}"
        p1.font.bold = True
        p1.font.size = Pt(11)
        p1.font.color.rgb = CYAN_ACCENT

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_LIGHT


    print("Building Slide 6: Tech Stack...")
    # ==========================================
    # SLIDE 6 — TECH STACK
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Technology Stack & Architectural Components", "TECH STACK", 6)

    stacks = [
        ("FRONTEND", [
            ("React 19", "Declarative component-based user interface"),
            ("Vite 8", "High-speed frontend development & bundling"),
            ("Tailwind CSS v4", "Custom dark-mode glassmorphic design system"),
            ("Recharts", "Interactive D3-based financial and telemetry charting"),
            ("React Router v7", "Single-page application client routing")
        ], CYAN_ACCENT),
        ("BACKEND", [
            ("Python 3.10+", "Core algorithmic backend language"),
            ("FastAPI", "High-performance asynchronous REST API framework"),
            ("Uvicorn", "Lightning-fast ASGI production web server"),
            ("Pydantic v2", "Strict schema validation and data contracts"),
            ("CORS Middleware", "Cross-origin resource sharing for frontend API calls")
        ], PRIMARY),
        ("DATA & MACHINE LEARNING", [
            ("Scikit-Learn", "Isolation Forest unsupervised anomaly detection & Linear Regression"),
            ("Pandas", "DataFrame manipulation, rolling statistics, and time-series aggregation"),
            ("NumPy", "High-performance vector operations and synthetic data generation"),
            ("SQLite3", "Self-contained serverless relational database engine")
        ], SUCCESS),
        ("CLOUD CONCEPTS (SIMULATED)", [
            ("AWS EC2", "Elastic Compute Cloud instances (m5, c5, t3, p3 families)"),
            ("AWS RDS", "Relational Database Service (PostgreSQL/MySQL db classes)"),
            ("AWS S3", "Simple Storage Service buckets (Standard vs Glacier/IA)"),
            ("AWS Lambda", "Serverless function invocations and concurrency telemetry")
        ], WARNING)
    ]

    for i, (cat_name, items, cat_color) in enumerate(stacks):
        col = i % 2
        row = i // 2
        left = Inches(0.8 + col * 5.95)
        top = Inches(1.6 + row * 2.6)
        width = Inches(5.75)
        height = Inches(2.45)

        create_card(s6, left, top, width, height, title=cat_name, border_color=cat_color)

        tb = s6.shapes.add_textbox(left + Inches(0.2), top + Inches(0.5), width - Inches(0.4), height - Inches(0.6))
        tf = tb.text_frame
        tf.word_wrap = True

        for tech, desc in items:
            add_bullet(tf, f"• {tech}:", desc, font_size=10)


    print("Building Slide 7: Novelty...")
    # ==========================================
    # SLIDE 7 — NOVELTY & KEY DIFFERENTIATORS
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Novelty & Key Differentiators — The Explainability Layer", "CORE NOVELTY", 7)

    # Top: Process Comparison
    create_card(s7, Inches(0.8), Inches(1.5), Inches(11.733), Inches(1.6), title="Traditional Cloud Monitoring vs CloudWise Approach")
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.333), Inches(1.1))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "Traditional Approach:", "Cloud Usage → Passive Billing Dashboard → Manual Human Triage → Opaque Decision (Reactive & Slow)", font_size=10.5, color=DANGER)
    add_bullet(tf, "CloudWise Approach:", "Cloud Usage → ML Anomaly Detection → Resource Heuristics → Explainable Recommendation → Quantified Savings (Proactive & Automated)", font_size=10.5, color=SUCCESS)

    # Bottom Left: 6 Key Differentiators
    create_card(s7, Inches(0.8), Inches(3.3), Inches(5.7), Inches(3.5), title="6 Key Platform Differentiators")
    tb = s7.shapes.add_textbox(Inches(1.0), Inches(3.75), Inches(5.3), Inches(2.9))
    tf = tb.text_frame
    tf.word_wrap = True

    diffs = [
        ("1. Explainable Recommendations:", "Answers WHAT is wrong, WHY it happened, and WHAT to do."),
        ("2. Machine Learning Anomaly Detection:", "Unsupervised Isolation Forest flags sudden spikes."),
        ("3. Deep Resource Intelligence:", "Combines billing data with CPU, connections, and request metrics."),
        ("4. Precise Savings Quantification:", "Every item is valued in Indian Rupees (₹) monthly & annually."),
        ("5. Predictive 14-Day Cost Forecasting:", "Linear trend projection alerts before monthly budget overruns."),
        ("6. Automated Executive Insights:", "CloudWise Insight engine generates plain-English takeaways.")
    ]
    for pfx, dsc in diffs:
        add_bullet(tf, pfx, dsc, font_size=9.5)

    # Bottom Right: Real Actual Recommendation Example
    create_card(s7, Inches(6.8), Inches(3.3), Inches(5.733), Inches(3.5), title="Actual System Recommendation Output", border_color=SUCCESS)
    tb2 = s7.shapes.add_textbox(Inches(7.0), Inches(3.75), Inches(5.333), Inches(2.9))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    add_bullet(tf2, "Flagged Resource:", "production-api-03 (EC2 m5.large in ap-south-1)", font_size=10.5, color=CYAN_ACCENT)
    add_bullet(tf2, "Identified Problem:", "EC2 instance is underutilized (Average CPU: 5.7%)", font_size=10.5, color=WARNING)
    add_bullet(tf2, "Current Cost:", "₹7,200 / month (~₹240 / day)", font_size=10.5)
    add_bullet(tf2, "Actionable Advice:", "Consider downsizing from m5.large to a smaller instance type (t3.medium) or schedule auto-shutdown during off-peak hours.", font_size=10)
    add_bullet(tf2, "Potential Savings:", "₹3,240 / month (₹38,880 / year)", font_size=11, color=SUCCESS)
    add_bullet(tf2, "Risk & Confidence:", "Risk: Low  |  Confidence Score: 92%", font_size=10.5, color=PRIMARY)
    add_bullet(tf2, "Why Was It Flagged?", '"CPU utilization averaged 5.7% over the 60-day observation period, well below the 10% threshold. The instance appears significantly oversized for its workload."', font_size=9.5, color=TEXT_MUTED)


    print("Building Slide 8: Initial Output Results...")
    # ==========================================
    # SLIDE 8 — INITIAL OUTPUT RESULTS (WITH SCREENSHOTS)
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Initial Output & Actual System Results", "EMPIRICAL RESULTS", 8)

    # Left: Actual KPIs & Summary Table
    create_card(s8, Inches(0.8), Inches(1.5), Inches(5.2), Inches(5.3), title="Actual Platform Metrics & KPIs")
    tb = s8.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(4.8), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "Total Simulated Cloud Spend:", "₹2,95,141 (~₹3.0 Lakhs across 60 days)", font_size=10.5)
    add_bullet(tf, "Potential Monthly Savings:", "₹16,310 / month (₹1,95,720 / year)", font_size=11, color=SUCCESS)
    add_bullet(tf, "Active Resources Monitored:", "20 assets (8 EC2, 4 RDS, 4 S3, 4 Lambda)", font_size=10.5)
    add_bullet(tf, "Anomalies Flagged:", "61 events detected (8 Critical, 2 High)", font_size=10.5, color=DANGER)
    add_bullet(tf, "Optimization Opportunities:", "8 active actionable recommendations", font_size=10.5, color=CYAN_ACCENT)
    add_bullet(tf, "Current Monthly Run-Rate:", "₹1,50,066 / month", font_size=10.5)
    add_bullet(tf, "Predicted Monthly Run-Rate:", "₹1,50,949 / month (+0.6% drift)", font_size=10.5)
    add_bullet(tf, "Forecast Prediction Window:", "14 days ahead", font_size=10.5)

    # Right: Embed Real Dashboard Screenshot
    dash_path = os.path.join(SCREENSHOT_DIR, "dashboard.png")
    if os.path.exists(dash_path):
        create_card(s8, Inches(6.2), Inches(1.5), Inches(6.333), Inches(5.3), title="Live Running Dashboard Screenshot")
        s8.shapes.add_picture(dash_path, Inches(6.35), Inches(2.0), width=Inches(6.033))
    else:
        create_card(s8, Inches(6.2), Inches(1.5), Inches(6.333), Inches(5.3), title="Live Running Dashboard")


    print("Building Slide 9: Implementation...")
    # ==========================================
    # SLIDE 9 — IMPLEMENTATION & PROGRAM OUTPUT
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Implementation & Program Source Code", "IMPLEMENTATION CODE", 9)

    # Left Column: Code Snippets
    create_card(s9, Inches(0.8), Inches(1.5), Inches(6.0), Inches(5.3), title="Core Backend Implementation (Python / FastAPI)")
    tb = s9.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.6), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "1. Isolation Forest Anomaly Detection (analysis.py):", "", font_size=10.5, color=CYAN_ACCENT)
    code1 = (
        'df["rolling_mean"] = df["daily_cost"].rolling(7, min_periods=1).mean()\n'
        'features = df[["daily_cost", "rolling_mean"]].values\n'
        'model = IsolationForest(contamination=0.08, random_state=42, n_estimators=40)\n'
        'preds = model.fit_predict(features)  # -1 flags anomalies'
    )
    add_bullet(tf, "", code1, font_size=8.5, color=TEXT_MUTED)
    add_bullet(tf, "One-Line Explanation:", "Unsupervised ML model isolates unusual spending deviations based on daily expenditure and 7-day rolling baseline.", font_size=9, color=TEXT_LIGHT)

    add_bullet(tf, "2. Underutilization & Dynamic Savings Rule (analysis.py):", "", font_size=10.5, color=CYAN_ACCENT)
    code2 = (
        'if r["service"] == "EC2" and r.get("utilization_status") == "Underutilized":\n'
        '    saving = round(r["monthly_cost"] * 0.45, 0)\n'
        '    confidence = min(95, round(90 + (10 - avg_cpu) * 0.5, 0))'
    )
    add_bullet(tf, "", code2, font_size=8.5, color=TEXT_MUTED)
    add_bullet(tf, "One-Line Explanation:", "Evaluates observed CPU heuristics against FinOps thresholds and dynamically scales confidence and savings.", font_size=9, color=TEXT_LIGHT)

    add_bullet(tf, "3. FastAPI Endpoint Handler (main.py):", "", font_size=10.5, color=CYAN_ACCENT)
    code3 = (
        '@app.get("/api/recommendations")\n'
        'def recommendations():\n'
        '    return analysis.get_recommendations()'
    )
    add_bullet(tf, "", code3, font_size=8.5, color=TEXT_MUTED)
    add_bullet(tf, "One-Line Explanation:", "Provides lightning-fast JSON payload serving recommendations directly to the React interface.", font_size=9, color=TEXT_LIGHT)

    # Right Column: Embed Real Anomaly Screenshot
    anom_path = os.path.join(SCREENSHOT_DIR, "anomalies.png")
    if os.path.exists(anom_path):
        create_card(s9, Inches(7.0), Inches(1.5), Inches(5.533), Inches(5.3), title="Live Anomaly Detection Output")
        s9.shapes.add_picture(anom_path, Inches(7.15), Inches(2.0), width=Inches(5.233))


    print("Building Slide 10: Graphs & Data Analysis...")
    # ==========================================
    # SLIDE 10 — GRAPHS & DATA ANALYSIS
    # ==========================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Graphs & Empirical Data Analysis", "DATA INSIGHTS", 10)

    # Top: Screenshot of Cost Analysis / Recommendations
    cost_path = os.path.join(SCREENSHOT_DIR, "cost_analysis.png")
    if os.path.exists(cost_path):
        create_card(s10, Inches(0.8), Inches(1.5), Inches(11.733), Inches(3.3), title="Real Visualizations Generated by CloudWise (Recharts / React)")
        s10.shapes.add_picture(cost_path, Inches(1.0), Inches(1.95), width=Inches(11.333))

    # Bottom: 4 Real Observations Supported by Data
    create_card(s10, Inches(0.8), Inches(4.95), Inches(11.733), Inches(1.9), title="Empirical Observations from Generated Telemetry")
    tb = s10.shapes.add_textbox(Inches(1.0), Inches(5.3), Inches(11.333), Inches(1.4))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "• Cost Distribution by Service:", "EC2 accounts for 48–51% (~₹1.50L) of cloud expenditure, RDS accounts for 40% (~₹1.18L), Lambda accounts for 6% (~₹17k), and S3 represents 3% (~₹10k).", font_size=10)
    add_bullet(tf, "• Daily Spending Trajectory:", "Daily spend maintains a consistent baseline between ₹4,700–₹5,200/day, with upward spikes occurring in the final 10 days due to injected Lambda traffic bursts and scaling events.", font_size=10)
    add_bullet(tf, "• Identified Inefficiencies:", "Three EC2 instances (production-api-03, staging-server, dev-server-02) maintain average CPU below 6%, representing ₹6,615/mo in preventable waste.", font_size=10)
    add_bullet(tf, "• 14-Day Linear Forecast:", "Linear regression indicates a mild upward trend (+0.6% monthly drift), alerting administrators to budget overruns before they materialize.", font_size=10)


    print("Building Slide 11: Scalability...")
    # ==========================================
    # SLIDE 11 — SCALABILITY & DEPLOYMENT POTENTIAL
    # ==========================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Scalability & Deployment Potential", "SYSTEM EVOLUTION", 11)

    # Top: Architecture Transition Diagram
    create_card(s11, Inches(0.8), Inches(1.5), Inches(11.733), Inches(2.2), title="Architecture Evolution: MVP to Production")
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(1.9), Inches(11.333), Inches(1.7))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "Current MVP Implementation:", "Simulated AWS Telemetry → SQLite Relational DB → FastAPI Engine → React Vite SPA (Self-contained, credential-free, fast demo)", font_size=10.5, color=CYAN_ACCENT)
    add_bullet(tf, "Production Enterprise Target:", "AWS Cost Explorer + CloudWatch Metric Streams + AWS Config → Scalable Ingestion Pipeline (Kafka / Celery) → PostgreSQL / TimescaleDB → CloudWise Analytics Core → React Enterprise Dashboard", font_size=10.5, color=SUCCESS)

    # Bottom: 8 Scalability Vectors (in 2 columns)
    create_card(s11, Inches(0.8), Inches(3.85), Inches(5.75), Inches(3.0), title="Infrastructure & Data Scalability")
    tb = s11.shapes.add_textbox(Inches(1.0), Inches(4.25), Inches(5.35), Inches(2.5))
    tf = tb.text_frame
    tf.word_wrap = True
    add_bullet(tf, "1. Database Evolution:", "Transition from SQLite to PostgreSQL / TimescaleDB for multi-year time-series telemetry.", font_size=9.5)
    add_bullet(tf, "2. Expanded Cloud Services:", "Broaden monitoring to AWS EKS (Kubernetes), ECS, DynamoDB, NAT Gateways, and CloudFront.", font_size=9.5)
    add_bullet(tf, "3. Scheduled Data Ingestion:", "Asynchronous event-driven data ingestion using AWS EventBridge, SQS, and Celery workers.", font_size=9.5)
    add_bullet(tf, "4. High-Throughput Processing:", "Distributed ML training and batched anomaly evaluation across millions of metrics.", font_size=9.5)

    create_card(s11, Inches(6.8), Inches(3.85), Inches(5.733), Inches(3.0), title="Enterprise Deployment & Security")
    tb2 = s11.shapes.add_textbox(Inches(7.0), Inches(4.25), Inches(5.333), Inches(2.5))
    tf2 = tb2.text_frame
    tf2.word_wrap = True
    add_bullet(tf2, "5. Live AWS Integration:", "Seamless connection via IAM Cross-Account Roles & AWS Boto3 SDK (no static keys).", font_size=9.5)
    add_bullet(tf2, "6. Multi-Account Management:", "Centralized FinOps visibility across entire AWS Organizations (Dev, Staging, Prod accounts).", font_size=9.5)
    add_bullet(tf2, "7. Security & Governance:", "Role-Based Access Control (RBAC), SSO integration, and full audit logging.", font_size=9.5)
    add_bullet(tf2, "8. Cloud-Native Hosting:", "Containerized Docker deployment on AWS ECS Fargate or AWS App Runner.", font_size=9.5)


    print("Building Slide 12: Future Scope...")
    # ==========================================
    # SLIDE 12 — FUTURE SCOPE
    # ==========================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Future Scope — Roadmap & Roadmap Distinctions", "FUTURE VISION", 12)

    # Left: Current MVP
    create_card(s12, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.3), title="Current MVP (Implemented & Confirmed)")
    tb = s12.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.3), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "✓ Simulated AWS Telemetry:", "60-day reproducible dataset covering EC2, RDS, S3, and Lambda instances with realistic noise and deliberate spikes.", font_size=10)
    add_bullet(tf, "✓ Unsupervised ML Anomaly Detection:", "Scikit-learn Isolation Forest flagging 61 cost anomalies with severity tiers (Critical, High, Medium, Low).", font_size=10)
    add_bullet(tf, "✓ Rule-Based FinOps Heuristics:", "Underutilization rules for compute (<10% CPU), database (<15% CPU, <10 conns), and storage access.", font_size=10)
    add_bullet(tf, "✓ Explainable Recommendation Layer:", "Dynamic answers for WHAT, WHY, ACTION, RISK, CONFIDENCE %, and SAVINGS in ₹ INR.", font_size=10)
    add_bullet(tf, "✓ 14-Day Cost Forecasting:", "Lightweight linear regression projecting short-term spending trends and percentage changes.", font_size=10)
    add_bullet(tf, "✓ Glassmorphic React Dashboard:", "Single-pane SaaS web interface with live Recharts visualizations and slide-out resource inspector.", font_size=10)

    # Right: Future Scope
    create_card(s12, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3), title="Future Development (Planned Enhancements)")
    tb2 = s12.shapes.add_textbox(Inches(7.0), Inches(1.95), Inches(5.333), Inches(4.7))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    add_bullet(tf2, "1. Live AWS Boto3 Integration:", "Direct credentialed connection to AWS Cost Explorer, CloudWatch Metrics, and AWS Compute Optimizer.", font_size=10)
    add_bullet(tf2, "2. Real-time Streaming Alarms:", "Integration with AWS SNS, Slack Webhooks, and PagerDuty for instantaneous cost spike alerts.", font_size=10)
    add_bullet(tf2, "3. Multi-Cloud FinOps Support:", "Extend intelligence layer to Microsoft Azure (Cost Management) and Google Cloud Platform (Cloud Billing).", font_size=10)
    add_bullet(tf2, "4. Advanced Time-Series ML:", "Evaluate SARIMA, Facebook Prophet, and LSTM neural networks for complex seasonal spending patterns.", font_size=10)
    add_bullet(tf2, "5. SHAP Feature Attribution:", "Integrate SHapley Additive exPlanations for mathematical feature-level explainability in enterprise audits.", font_size=10)
    add_bullet(tf2, "6. 1-Click Automated Remediation:", "Opt-in automated downsizing, schedule stop/start, and lifecycle rule migration with safety rollbacks.", font_size=10)


    print("Building Slide 13: Roadmap...")
    # ==========================================
    # SLIDE 13 — ROADMAP (NEXT 2 MONTHS - ACCELERATED)
    # ==========================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Project Roadmap — 2-Month Execution Timeline", "PROJECT ROADMAP", 13)

    # Top Banner: TARGET COMPLETION BY END OF MONTH 1
    banner = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(0.55))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(6, 78, 59)  # Deep Emerald
    banner.line.color.rgb = SUCCESS
    btf = banner.text_frame
    bp = btf.paragraphs[0]
    bp.text = "TARGET ACHIEVED: CORE PROJECT IMPLEMENTATION COMPLETED BY END OF MONTH 1"
    bp.alignment = PP_ALIGN.CENTER
    bp.font.size = Pt(11)
    bp.font.bold = True
    bp.font.color.rgb = RGBColor(167, 243, 208)

    # Month 1 Box (Left): Implementation & Wind-Up
    create_card(s13, Inches(0.8), Inches(2.2), Inches(5.7), Inches(4.6), title="Month 1: Implementation & Wind-Up (Done)", border_color=SUCCESS)
    tb1 = s13.shapes.add_textbox(Inches(1.0), Inches(2.65), Inches(5.3), Inches(4.0))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    add_bullet(tf1, "WEEK 1 (Foundation):", "Architecture design, simulated 60-day AWS dataset generation (EC2, RDS, S3, Lambda), SQLite database schema, and initial FastAPI REST server.", font_size=9.5)
    add_bullet(tf1, "WEEK 2 (Core Intelligence):", "Cost analysis engine, resource utilization telemetry calculations, and Scikit-learn Isolation Forest ML anomaly model.", font_size=9.5)
    add_bullet(tf1, "WEEK 3 (Explainability & Forecasting):", "Explainable recommendation engine, dynamic reasoning generator, 14-day linear cost forecasting, and CloudWise insights logic.", font_size=9.5)
    add_bullet(tf1, "WEEK 4 (Final Implementation / Wind-Up):", "Frontend-backend integration, Recharts visualizations, UI polish, layout bug fixes, endpoint testing, and final MVP completion.", font_size=9.5)

    # Month 2 Box (Right): Validation & Academic Deliverables
    create_card(s13, Inches(6.8), Inches(2.2), Inches(5.733), Inches(4.6), title="Month 2: Validation & Academic Deliverables", border_color=PRIMARY)
    tb2 = s13.shapes.add_textbox(Inches(7.0), Inches(2.65), Inches(5.333), Inches(4.0))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    add_bullet(tf2, "WEEK 5 (Validation & Stress Testing):", "Model hyperparameter tuning (contamination rates), boundary condition testing, and validation across varying resource loads.", font_size=9.5)
    add_bullet(tf2, "WEEK 6 (Comprehensive Documentation):", "Preparation of comprehensive project report, research documentation, architectural diagrams, and user guide.", font_size=9.5)
    add_bullet(tf2, "WEEK 7 (Presentation & Viva Rehearsal):", "Finalization of academic slide deck, live demo preparation, faculty review feedback incorporation, and practice dry runs.", font_size=9.5)
    add_bullet(tf2, "WEEK 8 (Final Evaluation & Submission):", "Final academic viva defense, project submission, GitHub code archive packaging, and evaluation sign-off.", font_size=9.5)


    print("Building Slide 14: Outcomes...")
    # ==========================================
    # SLIDE 14 — OUTCOMES (ACHIEVED VS FUTURE)
    # ==========================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Project Deliverables & Achieved Outcomes", "DELIVERABLES", 14)

    # Left: Currently Achieved
    create_card(s14, Inches(0.8), Inches(1.5), Inches(5.7), Inches(5.3), title="Currently Achieved Deliverables (Verified)", border_color=SUCCESS)
    tb = s14.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(5.3), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "✓ Fully Functional FinOps Dashboard:", "High-performance React 19 web application with dynamic KPIs, dark-mode design system, and zero layout issues.", font_size=10)
    add_bullet(tf, "✓ Reproducible 60-Day AWS Dataset:", "20 simulated cloud assets (EC2, RDS, S3, Lambda) with realistic seasonal variance and injected anomalies.", font_size=10)
    add_bullet(tf, "✓ Unsupervised ML Cost Anomaly Detection:", "Isolation Forest algorithm successfully identifying 61 cost spikes categorized into 4 severity levels.", font_size=10)
    add_bullet(tf, "✓ Transparent Utilization Profiling:", "Accurate detection of underutilized instances, idle databases, and infrequently accessed storage.", font_size=10)
    add_bullet(tf, "✓ Explainable Recommendation Engine:", "Every opportunity provides Problem, Action, Confidence %, Risk tier, and human-readable Reason.", font_size=10)
    add_bullet(tf, "✓ Demonstrated Financial Savings:", "Identified ₹16,310/month (~₹1.95 Lakhs/year) of potential recurring savings across 8 resources.", font_size=10)
    add_bullet(tf, "✓ 14-Day Predictive Spending Forecast:", "Linear regression model forecasting upcoming expenditure trends (+0.6% monthly drift).", font_size=10)

    # Right: Future Expected Outcomes
    create_card(s14, Inches(6.8), Inches(1.5), Inches(5.733), Inches(5.3), title="Future / Expected Outcomes (Planned)", border_color=CARD_BORDER)
    tb2 = s14.shapes.add_textbox(Inches(7.0), Inches(1.95), Inches(5.333), Inches(4.7))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    add_bullet(tf2, "• Live AWS API Integration:", "Real-time synchronization with AWS Cost Explorer, CloudWatch Metric Streams, and AWS Budgets via IAM.", font_size=10)
    add_bullet(tf2, "• Multi-Cloud Support:", "Unified cost intelligence dashboard bridging AWS, Microsoft Azure, and Google Cloud Platform.", font_size=10)
    add_bullet(tf2, "• Automated 1-Click Remediation:", "Direct execution of recommended downsizing, stopping idle instances, or transitioning S3 storage tiers.", font_size=10)
    add_bullet(tf2, "• Enterprise Organization Management:", "Multi-tenant role-based access control, SSO authentication, and team-level budget alerting.", font_size=10)
    add_bullet(tf2, "• Continuous Real-Time Streaming Anomaly Detection:", "Kafka/Kinesis streaming anomaly detection for instantaneous alerts.", font_size=10)


    print("Building Slide 15: References...")
    # ==========================================
    # SLIDE 15 — REFERENCES
    # ==========================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Academic & Technical References", "REFERENCES", 15)

    create_card(s15, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.3), title="Official Documentation & Academic Literature")
    tb = s15.shapes.add_textbox(Inches(1.1), Inches(1.95), Inches(11.133), Inches(4.7))
    tf = tb.text_frame
    tf.word_wrap = True

    refs = [
        ("1. AWS Cost Management Documentation (2024):", "Amazon Web Services. 'AWS Cost Explorer and AWS Cost Anomaly Detection User Guide.' AWS Cloud Financial Management Documentation, Seattle, WA."),
        ("2. AWS CloudWatch Metrics Reference (2024):", "Amazon Web Services. 'Amazon CloudWatch Metrics, Dimensions, and Telemetry Reference Guide.' AWS Technical Documentation."),
        ("3. Isolation Forest Algorithm:", "Liu, F. T., Ting, K. M., & Zhou, Z. H. (2008). 'Isolation Forest.' In Proceedings of the Eighth IEEE International Conference on Data Mining (ICDM), pp. 413–422. DOI: 10.1109/ICDM.2008.17."),
        ("4. Scikit-Learn Framework:", "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., et al. (2011). 'Scikit-learn: Machine Learning in Python.' Journal of Machine Learning Research (JMLR), 12, pp. 2825–2830."),
        ("5. Cloud FinOps Best Practices:", "Storment, J. R., & Fuller, M. (2023). 'Cloud FinOps: Collaborative, Real-Time Cloud Financial Management.' 2nd Edition, O'Reilly Media, Sebastopol, CA."),
        ("6. FastAPI High-Performance Framework:", "Tiangolo, S. (2023). 'FastAPI: Modern, Fast (High-Performance) Web Framework for Building APIs with Python 3.8+.' https://fastapi.tiangolo.com"),
        ("7. React User Interface Library:", "Meta Open Source (2024). 'React: The Library for Web and Native User Interfaces.' https://react.dev"),
        ("8. Recharts Data Visualization:", "Recharts Group (2024). 'Recharts: A Composable Charting Library Built on React Components.' https://recharts.org")
    ]

    for pfx, cit in refs:
        add_bullet(tf, pfx, cit, font_size=10.5)


    print("Building Slide 16: Conclusion...")
    # ==========================================
    # SLIDE 16 — CONCLUSION
    # ==========================================
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "Conclusion & Project Summary", "CONCLUSION", 16)

    # Top Card: Summary
    create_card(s16, Inches(0.8), Inches(1.5), Inches(11.733), Inches(3.2), title="Project Summary & Key Takeaways")
    tb = s16.shapes.add_textbox(Inches(1.0), Inches(1.95), Inches(11.333), Inches(2.6))
    tf = tb.text_frame
    tf.word_wrap = True

    add_bullet(tf, "• Transformation of Cloud Billing:", "CloudWise demonstrates how raw cloud cost data can be transformed from passive retrospective billing reports into actionable, explainable resource intelligence.", font_size=11)
    add_bullet(tf, "• Bridging AI and FinOps Operations:", "By coupling unsupervised Isolation Forest anomaly detection with deterministic telemetry heuristics, the system eliminates bill shock without requiring manual threshold configuration.", font_size=11)
    add_bullet(tf, "• Quantifiable Financial Impact:", "Successfully demonstrated ₹16,310/month (~₹1.95 Lakhs/year) of potential recurring savings across 8 resources, complete with confidence scores and non-destructive recommendations.", font_size=11)
    add_bullet(tf, "• Transparent Explainability:", "Engineers and executives can confidently make rightsizing decisions because every recommendation answers WHAT is wrong, WHY it was flagged, and WHAT to do.", font_size=11)

    # Bottom Creed Card: Core Creed + Thank You
    b_card = s16.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.9), Inches(11.733), Inches(1.9))
    b_card.fill.solid()
    b_card.fill.fore_color.rgb = RGBColor(19, 29, 53)
    b_card.line.color.rgb = PRIMARY
    btf = b_card.text_frame
    btf.word_wrap = True

    bp1 = btf.paragraphs[0]
    bp1.text = '"From cloud expenditure → to cloud intelligence → to informed optimization."'
    bp1.alignment = PP_ALIGN.CENTER
    bp1.font.size = Pt(14)
    bp1.font.bold = True
    bp1.font.color.rgb = CYAN_ACCENT

    bp2 = btf.add_paragraph()
    bp2.text = "THANK YOU"
    bp2.alignment = PP_ALIGN.CENTER
    bp2.font.size = Pt(28)
    bp2.font.bold = True
    bp2.font.color.rgb = TEXT_LIGHT

    bp3 = btf.add_paragraph()
    bp3.text = "Questions & Faculty Discussion | Pragya Prakash (24BML0075) & Vidhi Agarwal (24BML0081) | Guide: Prof. Sangeetha N"
    bp3.alignment = PP_ALIGN.CENTER
    bp3.font.size = Pt(10)
    bp3.font.color.rgb = TEXT_MUTED

    print(f"Saving presentation to {OUTPUT_FILE}...")
    prs.save(OUTPUT_FILE)
    print("Presentation successfully created and saved!")


if __name__ == "__main__":
    build_all_slides()
